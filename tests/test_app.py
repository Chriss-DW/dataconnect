"""Pruebas de integración con Excel temporal; nunca usan contactos personales."""
from concurrent.futures import ThreadPoolExecutor
from pathlib import Path
import tempfile
import unittest
from unittest.mock import patch

from openpyxl import load_workbook

from dataconnect import create_app
from dataconnect.routes import safe_next
from dataconnect.storage import ContactStore, HEADERS, StorageError


class DataConnectTests(unittest.TestCase):
    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.filename = Path(self.directory.name) / "contacts.xlsx"
        self.app = create_app({"TESTING": True, "DEMO_MODE": False,
                               "SECRET_KEY": "a-test-secret-that-is-not-used-in-production",
                               "ADMIN_USERNAME": "tester", "ADMIN_PASSWORD": "only-for-testing-123",
                               "DATA_FILE": self.filename})
        self.client = self.app.test_client()
        self.store = self.app.extensions["contacts"]

    def tearDown(self):
        self.directory.cleanup()

    def token(self, client=None):
        client = client or self.client
        client.get("/login")
        with client.session_transaction() as state:
            return state["csrf_token"]

    def login(self):
        return self.client.post("/login", data={"username": "tester", "password": "only-for-testing-123",
                                                 "csrf_token": self.token()})

    def post(self, url, data=None):
        with self.client.session_transaction() as state:
            token = state["csrf_token"]
        return self.client.post(url, data={**(data or {}), "csrf_token": token})

    def payload(self, **changes):
        return {"full_name": "Ana Prueba", "phone": "00000001", "email": "ana@example.test",
                "address": "Dirección ficticia", "category": "Trabajo", **changes}

    def add(self, **changes):
        values = self.payload(**changes)
        parts = values.pop("full_name").split(maxsplit=1)
        values.update(first_name=parts[0], last_name=parts[1] if len(parts) > 1 else "")
        return self.store.create(values)

    def test_root_redirects_to_login(self):
        self.assertEqual(self.client.get("/").location, "/login")

    def test_all_private_views_require_login(self):
        contact_id = self.add()
        for url in ("/contactos", "/contactos/nuevo", "/estadisticas", f"/contactos/{contact_id}",
                    f"/contactos/{contact_id}/editar", f"/contactos/{contact_id}/eliminar"):
            with self.subTest(url=url):
                response = self.client.get(url)
                self.assertEqual(response.status_code, 302)
                self.assertIn("/login?next=", response.location)

    def test_unauthenticated_posts_cannot_mutate(self):
        contact_id = self.add()
        token = self.token()
        for url in ("/contactos/nuevo", f"/contactos/{contact_id}/editar",
                    f"/contactos/{contact_id}/eliminar", f"/contactos/{contact_id}/favorito"):
            response = self.client.post(url, data={**self.payload(), "csrf_token": token})
            self.assertEqual(response.status_code, 302)
        self.assertEqual(len(self.store.all()), 1)
        self.assertFalse(self.store.get(contact_id)["favorite"])

    def test_login_creates_session_and_rotates_csrf(self):
        original = self.token()
        response = self.login()
        self.assertEqual(response.status_code, 302)
        with self.client.session_transaction() as state:
            self.assertTrue(state["authenticated"])
            self.assertNotEqual(state["csrf_token"], original)

    def test_wrong_password_is_visible_and_preserves_username(self):
        response = self.client.post("/login", data={"username": "tester", "password": "wrong", "csrf_token": self.token()})
        self.assertEqual(response.status_code, 422)
        self.assertIn(b'value="tester"', response.data)
        self.assertIn("no son correctos".encode(), response.data)

    def test_csrf_is_required_even_for_login(self):
        response = self.client.post("/login", data={"username": "tester", "password": "only-for-testing-123"})
        self.assertEqual(response.status_code, 400)

    def test_non_ascii_csrf_is_rejected_without_server_error(self):
        self.assertEqual(self.client.post("/login", data={"csrf_token": "ñ"}).status_code, 400)

    def test_logout_clears_session(self):
        self.login()
        self.assertEqual(self.post("/logout").status_code, 302)
        self.assertEqual(self.client.get("/contactos").status_code, 302)

    def test_logout_get_is_not_allowed(self):
        self.assertEqual(self.client.get("/logout").status_code, 405)

    def test_external_next_is_rejected(self):
        for destination in ("https://example.test", "//example.test", "/\\example.test", "/\nexample.test"):
            self.assertEqual(safe_next(destination, "/contactos"), "/contactos")
        self.assertEqual(safe_next("/estadisticas", "/contactos"), "/estadisticas")

    def test_login_can_return_to_an_internal_page(self):
        response = self.client.post("/login?next=/estadisticas", data={"username": "tester", "password": "only-for-testing-123", "csrf_token": self.token()})
        self.assertEqual(response.location, "/estadisticas")

    def test_create_contact_persists_every_field(self):
        self.login()
        response = self.post("/contactos/nuevo", self.payload())
        self.assertEqual(response.status_code, 302)
        contact = self.store.all()[0]
        self.assertEqual(contact["full_name"], "Ana Prueba")
        self.assertEqual(contact["phone"], "00000001")
        self.assertEqual(contact["email"], "ana@example.test")
        self.assertEqual(contact["address"], "Dirección ficticia")
        self.assertEqual(contact["category"], "Trabajo")
        self.assertFalse(contact["favorite"])

    def test_phone_requires_eight_ascii_digits(self):
        self.login()
        for phone in ("123", "0000000a", "１２３４５６７８", "000000001", "1234 567"):
            with self.subTest(phone=phone):
                self.assertEqual(self.post("/contactos/nuevo", self.payload(phone=phone)).status_code, 422)
        self.assertEqual(self.store.all(), [])

    def test_long_email_domain_is_valid(self):
        self.login()
        self.assertEqual(self.post("/contactos/nuevo", self.payload(email="ana@example.technology")).status_code, 302)

    def test_invalid_email_is_rejected(self):
        self.login()
        for email in ("", "ana", "ana@", "ana@example", "ana x@example.test"):
            self.assertEqual(self.post("/contactos/nuevo", self.payload(email=email)).status_code, 422)

    def test_address_is_optional(self):
        self.login()
        self.assertEqual(self.post("/contactos/nuevo", self.payload(address="")).status_code, 302)

    def test_invalid_category_is_rejected(self):
        self.login()
        self.assertEqual(self.post("/contactos/nuevo", self.payload(category="Inventada")).status_code, 422)

    def test_error_preserves_input_and_explains_the_field(self):
        self.login()
        response = self.post("/contactos/nuevo", self.payload(full_name="Nombre que se conserva", phone="12"))
        self.assertEqual(response.status_code, 422)
        self.assertIn(b'Nombre que se conserva', response.data)
        self.assertIn(b'id="error-phone"', response.data)
        self.assertIn(b'aria-invalid="true"', response.data)

    def test_limits_and_illegal_excel_characters_are_validated(self):
        self.login()
        for changes in ({"full_name": "a" * 121}, {"address": "a" * 241}, {"full_name": "Ana\x00Prueba"}):
            self.assertEqual(self.post("/contactos/nuevo", self.payload(**changes)).status_code, 422)
        self.assertEqual(self.store.all(), [])

    def test_edit_preserves_id_and_favorite(self):
        contact_id = self.add()
        self.store.update(contact_id, toggle=True)
        self.login()
        response = self.post(f"/contactos/{contact_id}/editar", self.payload(full_name="Ana Actualizada"))
        self.assertEqual(response.status_code, 302)
        contact = self.store.get(contact_id)
        self.assertEqual(contact["full_name"], "Ana Actualizada")
        self.assertTrue(contact["favorite"])
        self.assertEqual(contact["id"], contact_id)

    def test_invalid_edit_does_not_change_the_record(self):
        contact_id = self.add()
        original = self.store.get(contact_id)
        self.login()
        self.assertEqual(self.post(f"/contactos/{contact_id}/editar", self.payload(phone="bad")).status_code, 422)
        self.assertEqual(self.store.get(contact_id), original)

    def test_get_delete_only_shows_confirmation(self):
        contact_id = self.add()
        self.login()
        self.assertEqual(self.client.get(f"/contactos/{contact_id}/eliminar").status_code, 200)
        self.assertIsNotNone(self.store.get(contact_id))

    def test_delete_is_post_and_checks_csrf(self):
        contact_id = self.add()
        self.login()
        self.assertEqual(self.client.post(f"/contactos/{contact_id}/eliminar").status_code, 400)
        self.assertIsNotNone(self.store.get(contact_id))
        self.assertEqual(self.post(f"/contactos/{contact_id}/eliminar").status_code, 302)
        self.assertIsNone(self.store.get(contact_id))

    def test_deleting_a_row_does_not_change_another_contact_id(self):
        first = self.add()
        second = self.add(full_name="Luis Ejemplo", email="luis@example.test")
        third = self.add(full_name="Lucía Muestra", email="lucia@example.test")
        self.store.delete(first)
        self.login()
        self.post(f"/contactos/{second}/editar", self.payload(full_name="Luis Editado"))
        self.assertEqual(self.store.get(second)["full_name"], "Luis Editado")
        self.assertEqual(self.store.get(third)["full_name"], "Lucía Muestra")

    def test_invalid_ids_do_not_modify_headers_or_create_gaps(self):
        self.add()
        self.login()
        for contact_id in ("0", "1", "20", "missing"):
            for suffix in ("", "/editar", "/eliminar"):
                self.assertEqual(self.client.get(f"/contactos/{contact_id}{suffix}").status_code, 404)
            self.assertEqual(self.post(f"/contactos/{contact_id}/editar", self.payload()).status_code, 404)
        book = load_workbook(self.filename)
        self.assertEqual([cell.value for cell in book.active[1]], HEADERS)
        self.assertEqual(book.active.max_row, 2)
        book.close()

    def test_favorite_can_be_toggled_and_filtered(self):
        contact_id = self.add()
        self.login()
        self.assertEqual(self.client.get(f"/contactos/{contact_id}/favorito").status_code, 405)
        self.post(f"/contactos/{contact_id}/favorito")
        self.assertTrue(self.store.get(contact_id)["favorite"])
        self.assertIn(b'Ana Prueba', self.client.get("/contactos?favoritos=1").data)
        self.post(f"/contactos/{contact_id}/favorito")
        self.assertFalse(self.store.get(contact_id)["favorite"])

    def test_search_supports_accents_and_single_phone_digit(self):
        self.add(full_name="Lucía Ejemplo", phone="00000007")
        self.login()
        for query in ("lucia", "LUCÍA", "7", "example.test", "ficticia"):
            self.assertIn("Lucía Ejemplo".encode(), self.client.get("/contactos", query_string={"q": query}).data)

    def test_sorting_and_category_filter(self):
        self.add(full_name="Zoe Prueba", category="Amigos")
        self.add(full_name="Ana Prueba", category="Trabajo")
        self.login()
        response = self.client.get("/contactos?orden=nombre").data
        self.assertLess(response.index(b'Ana Prueba'), response.index(b'Zoe Prueba'))
        filtered = self.client.get("/contactos?categoria=Amigos").data
        self.assertIn(b'Zoe Prueba', filtered)
        self.assertNotIn(b'Ana Prueba', filtered)

    def test_empty_directory_and_no_results_are_distinct(self):
        self.login()
        self.assertIn("Tu próxima conexión empieza aquí".encode(), self.client.get("/contactos").data)
        self.assertIn(b'No encontramos coincidencias', self.client.get("/contactos?q=missing").data)

    def test_statistics_derive_from_stored_data(self):
        first = self.add()
        self.add(full_name="Luis Ejemplo", category="Familia")
        self.store.update(first, toggle=True)
        self.login()
        response = self.client.get("/estadisticas")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b'<dd>2</dd>', response.data)
        self.assertIn(b'<dd>1</dd>', response.data)
        self.assertNotIn(b'exponencial', response.data)
        self.assertNotIn(b'Conectado', response.data)

    def test_formulas_are_stored_as_literal_text(self):
        self.add(full_name="=1+1 Prueba", address="=2+2")
        book = load_workbook(self.filename)
        self.assertEqual(book.active.cell(2, 2).data_type, "s")
        self.assertEqual(book.active.cell(2, 6).data_type, "s")
        self.assertEqual(book.active.cell(2, 6).value, "=2+2")
        book.close()

    def test_html_is_escaped(self):
        self.add(full_name="<script>alert(1)</script>")
        self.login()
        response = self.client.get("/contactos")
        self.assertNotIn(b'<script>alert(1)</script>', response.data)
        self.assertIn(b'&lt;script&gt;', response.data)

    def test_simultaneous_writes_do_not_lose_contacts(self):
        def add(index):
            return self.add(full_name=f"Persona Prueba{index}", phone=f"{index:08d}")
        with ThreadPoolExecutor(max_workers=8) as executor:
            identifiers = list(executor.map(add, range(20)))
        self.assertEqual(len(self.store.all()), 20)
        self.assertEqual(len(set(identifiers)), 20)

    def test_failed_atomic_replace_keeps_original_workbook(self):
        self.add()
        original = self.filename.read_bytes()
        with patch.object(Path, "replace", side_effect=PermissionError("test lock")):
            with self.assertRaises(StorageError):
                self.add(full_name="No se debe guardar")
        self.assertEqual(self.filename.read_bytes(), original)
        self.assertEqual(len(self.store.all()), 1)
        self.assertEqual(list(self.filename.parent.glob("tmp*.xlsx")), [])

    def test_corrupt_or_missing_book_does_not_silently_reset_data(self):
        self.login()
        self.filename.write_bytes(b"not a workbook")
        response = self.client.get("/contactos")
        self.assertEqual(response.status_code, 503)
        self.assertNotIn(b'Traceback', response.data)
        self.assertEqual(self.filename.read_bytes(), b"not a workbook")
        self.filename.unlink()
        self.assertEqual(self.client.get("/contactos").status_code, 503)
        self.assertFalse(self.filename.exists())

    def test_format_mismatch_is_not_overwritten(self):
        book = load_workbook(self.filename)
        book.active.cell(1, 1).value = "Incorrecto"
        book.save(self.filename)
        book.close()
        before = self.filename.read_bytes()
        with self.assertRaises(StorageError):
            ContactStore(self.filename)
        self.assertEqual(self.filename.read_bytes(), before)

    def test_session_cookie_and_security_headers(self):
        response = self.client.get("/login")
        self.assertIn("HttpOnly", response.headers["Set-Cookie"])
        self.assertIn("SameSite=Lax", response.headers["Set-Cookie"])
        self.assertEqual(response.headers["X-Frame-Options"], "DENY")
        self.assertEqual(response.headers["X-Content-Type-Options"], "nosniff")
        self.assertIn("script-src 'self'", response.headers["Content-Security-Policy"])
        self.assertEqual(response.headers["Cache-Control"], "no-store")

    def test_password_is_kept_as_hash_in_configuration(self):
        self.assertNotIn("ADMIN_PASSWORD", self.app.config)
        self.assertTrue(self.app.config["PASSWORD_HASH"].startswith("scrypt:"))

    def test_private_mode_requires_non_demo_configuration(self):
        with self.assertRaises(ValueError):
            create_app({"DEMO_MODE": False, "SECRET_KEY": "short", "ADMIN_USERNAME": "x", "ADMIN_PASSWORD": "short"})

    def test_large_form_is_rejected(self):
        self.assertEqual(self.client.post("/login", data={"username": "x" * 70000}).status_code, 413)

    def test_demo_seed_is_fictitious_and_idempotent(self):
        self.store.seed_demo()
        self.store.seed_demo()
        self.assertEqual(len(self.store.all()), 6)
        self.assertTrue(all(contact["email"].endswith("@example.test") for contact in self.store.all()))
        self.assertEqual(sum(contact["favorite"] for contact in self.store.all()), 2)

    def test_readback_survives_a_new_store_instance(self):
        contact_id = self.add()
        reopened = ContactStore(self.filename)
        self.assertEqual(reopened.get(contact_id)["full_name"], "Ana Prueba")


if __name__ == "__main__":
    unittest.main(verbosity=2)
