"""Persistencia Excel con identificadores estables y escritura atómica.

El bloqueo coordina los hilos de una sola instancia de Python. Para varios
procesos o servidores se debe sustituir Excel por una base de datos transaccional.
"""
from __future__ import annotations

from contextlib import contextmanager
from datetime import datetime, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from threading import RLock
from uuid import uuid4

from openpyxl import Workbook, load_workbook

HEADERS = ["ID", "Nombre", "Apellido", "Teléfono", "Correo", "Dirección",
           "Categoría", "Favorito", "Creado", "Actualizado"]
FIELDS = ["id", "first_name", "last_name", "phone", "email", "address",
          "category", "favorite", "created_at", "updated_at"]
CATEGORIES = ("Trabajo", "Familia", "Amigos", "Otros")


class StorageError(RuntimeError):
    """El archivo no se pudo leer o guardar sin comprometer sus datos."""


class ContactStore:
    def __init__(self, filename: str | Path):
        self.path = Path(filename).resolve()
        self._lock = RLock()
        self.path.parent.mkdir(parents=True, exist_ok=True)
        with self._lock:
            if not self.path.exists():
                book = Workbook()
                book.active.title = "Contactos"
                book.active.append(HEADERS)
                try:
                    self._save(book)
                finally:
                    book.close()
            self.all()

    @contextmanager
    def _workbook(self):
        book = None
        try:
            book = load_workbook(self.path)
            if [cell.value for cell in book.active[1]] != HEADERS:
                raise StorageError("El archivo de contactos no tiene el formato esperado.")
            yield book
        except StorageError:
            raise
        except (OSError, ValueError, KeyError, EOFError) as error:
            raise StorageError("No fue posible acceder al archivo de contactos.") from error
        except Exception as error:
            # Un libro dañado no debe producir una página con detalles internos.
            raise StorageError("El archivo de contactos no se pudo procesar.") from error
        finally:
            if book is not None:
                book.close()

    def _save(self, book):
        temporary = None
        try:
            with NamedTemporaryFile(dir=self.path.parent, suffix=".xlsx", delete=False) as handle:
                temporary = Path(handle.name)
            book.save(temporary)
            temporary.replace(self.path)
        except OSError as error:
            raise StorageError("No se pudo guardar. Cerrá el archivo Excel e intentá de nuevo.") from error
        finally:
            if temporary is not None:
                temporary.unlink(missing_ok=True)

    @staticmethod
    def _record(row):
        contact = dict(zip(FIELDS, row))
        contact = {key: ("" if value is None else str(value)) for key, value in contact.items()}
        contact["favorite"] = contact["favorite"] == "Sí"
        contact["full_name"] = f'{contact["first_name"]} {contact["last_name"]}'.strip()
        contact["initials"] = "".join(part[0] for part in contact["full_name"].split()[:2]).upper()
        return contact

    @staticmethod
    def _write_row(sheet, number, contact):
        for column, field in enumerate(FIELDS, start=1):
            value = contact[field]
            if field == "favorite":
                value = "Sí" if value else "No"
            cell = sheet.cell(row=number, column=column, value=str(value))
            # Entradas que comienzan con '=' se conservan como texto, no fórmulas.
            cell.data_type = "s"

    def all(self):
        with self._lock, self._workbook() as book:
            return [self._record(row) for row in book.active.iter_rows(min_row=2, values_only=True) if row[0]]

    def get(self, contact_id):
        return next((contact for contact in self.all() if contact["id"] == contact_id), None)

    def create(self, values):
        now = datetime.now(timezone.utc).isoformat()
        contact = {**values, "id": uuid4().hex, "favorite": False,
                   "created_at": now, "updated_at": now}
        with self._lock, self._workbook() as book:
            self._write_row(book.active, book.active.max_row + 1, contact)
            self._save(book)
        return contact["id"]

    def update(self, contact_id, values=None, toggle=False):
        with self._lock, self._workbook() as book:
            for number, row in enumerate(book.active.iter_rows(min_row=2, values_only=True), start=2):
                if row[0] != contact_id:
                    continue
                contact = self._record(row)
                contact.update(values or {})
                if toggle:
                    contact["favorite"] = not contact["favorite"]
                contact["updated_at"] = datetime.now(timezone.utc).isoformat()
                self._write_row(book.active, number, contact)
                self._save(book)
                return True
            return False

    def delete(self, contact_id):
        with self._lock, self._workbook() as book:
            for number, row in enumerate(book.active.iter_rows(min_row=2, values_only=True), start=2):
                if row[0] == contact_id:
                    book.active.delete_rows(number)
                    self._save(book)
                    return True
            return False

    def seed_demo(self):
        if self.all():
            return
        examples = [
            ("Ana", "Prueba", "ana", "Trabajo"), ("Mateo", "Ejemplo", "mateo", "Amigos"),
            ("Lucía", "Muestra", "lucia", "Familia"), ("Diego", "Prueba", "diego", "Trabajo"),
            ("Valeria", "Ejemplo", "valeria", "Otros"), ("Nicolás", "Muestra", "nicolas", "Amigos"),
        ]
        for index, (first, last, email, category) in enumerate(examples, start=1):
            contact_id = self.create({"first_name": first, "last_name": last,
                                      "phone": f"000000{index:02d}", "email": f"{email}@example.test",
                                      "address": "Dirección ficticia para demostración", "category": category})
            if index in (1, 3):
                self.update(contact_id, toggle=True)
