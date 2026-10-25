"""Independent characterization tests. Only disposable workbooks are used."""
import os, sys, json, tempfile, importlib.util, contextlib, io, threading
from pathlib import Path
from importlib.metadata import version
from openpyxl import Workbook, load_workbook
from flask import template_rendered

ROOT = Path(__file__).resolve().parent
sys.dont_write_bytecode = True
SOURCE = ROOT / 'original-app' / 'app.py'
results = []

def record(name, observed):
    results.append({'test': name, 'observed': observed})

with tempfile.TemporaryDirectory(prefix='dataconnect-audit-') as temp:
    os.chdir(temp)  # import-time initialization cannot touch original/shared Excel
    spec = importlib.util.spec_from_file_location('isolated_dataconnect', SOURCE)
    module = importlib.util.module_from_spec(spec)
    sys.modules[spec.name] = module
    spec.loader.exec_module(module)
    module.EXCEL_FILE = str(Path(temp) / 'contactos.xlsx')
    app = module.app
    app.config.update(TESTING=False)
    contexts = []
    def capture(sender, template, context, **extra):
        contexts.append({'template': template.name, 'context': context})
    template_rendered.connect(capture, app)
    headers = ['Nombre','Apellido','Teléfono','Correo','Dirección','Categoría','Favorito']
    def reset(rows=None):
        wb = Workbook(); ws = wb.active; ws.append(headers)
        for row in rows or []: ws.append(row)
        wb.save(module.EXCEL_FILE); wb.close()
    def rows():
        wb=load_workbook(module.EXCEL_FILE); data=list(wb.active.values); wb.close(); return data
    a=['Ana','Prueba','00000001','ana@example.test','Centro','Familia','Sí']
    b=['Luis','Ejemplo','00000002','luis@example.test','Norte','Trabajo','No']
    c=['Zoe','Muestra','00000003','zoe@example.test','Sur','Amigos','No']
    valid={'nombre_completo':'Caso Temporal','telefono':'12345678','correo':'caso@example.test','direccion':'Centro','categoria':'Familia'}
    client=app.test_client()
    reset([a,b,c])
    record('anonymous_reads', {p:client.get(p).status_code for p in ['/principal','/crear_contacto','/editar/2','/mas_informacion/2','/estadisticas']})
    response=client.post('/login',data={'usuario':'admin','contrasena':'1234'})
    record('login', {'status':response.status_code,'location':response.location,'set_cookie':response.headers.get('Set-Cookie')})
    client.get('/login')
    record('after_logout_link', client.get('/principal').status_code)
    reset([a,b])
    response=client.post('/guardar_contacto',data=valid)
    record('anonymous_create_no_csrf',{'status':response.status_code,'location':response.location,'rows':len(rows())-1})
    response=client.get('/toggle_favorito/2')
    record('anonymous_get_favorite',{'status':response.status_code,'value':rows()[1][6]})
    response=client.get('/eliminar_contacto/2')
    record('anonymous_get_delete',{'status':response.status_code,'first_name_after':rows()[1][0]})
    reset([a,b,c]); client.get('/editar/3'); client.get('/eliminar_contacto/2')
    edit={**valid,'nombre':'Luis Editado'}; edit.pop('nombre_completo')
    response=client.post('/actualizar_contacto/3',data=edit)
    record('stale_edit_id_after_delete',{'status':response.status_code,'remaining_names':[r[0:2] for r in rows()[1:]],'lost_contact':'Zoe Muestra'})
    reset([a,b]); response=client.post('/actualizar_contacto/1',data=edit)
    record('update_header_id_1',{'status':response.status_code,'header':rows()[0]})
    reset([a]); response=client.post('/actualizar_contacto/20',data=edit)
    record('update_nonexistent_id_20',{'status':response.status_code,'row_count':len(rows()),'row20':rows()[19]})
    reset([a]); response=client.post('/actualizar_contacto/0',data=edit)
    record('update_id_0',{'status':response.status_code})
    reset([])
    for label,changes in [('missing_direction_category',{'direccion':'','categoria':'Inventada'}),('unicode_digits',{'telefono':'１２３４５６７８'}),('formula',{'nombre_completo':'=1+1','direccion':'=2+2'}),('long_tld',{'correo':'caso@example.technology'}),('bad_phone',{'telefono':'123'}),('bad_email',{'correo':'invalido'})]:
        before=len(rows()); response=client.post('/guardar_contacto',data={**valid,**changes})
        record('validation_'+label,{'status':response.status_code,'location':response.location,'row_delta':len(rows())-before,'error_visible_in_redirect':'Error' in client.get(response.location).get_data(as_text=True)})
    wb=load_workbook(module.EXCEL_FILE)
    record('formula_cell_types',{'name':wb.active['A4'].data_type,'direction':wb.active['E4'].data_type}); wb.close()
    reset([a,b,c]); client.get('/principal?q=0'); record('search_single_phone_digit',len(contexts[-1]['context']['contactos']))
    client.get('/principal?q=00000001'); record('search_full_phone',len(contexts[-1]['context']['contactos']))
    client.get('/principal?q=Familia'); record('search_category',len(contexts[-1]['context']['contactos']))
    client.get('/principal?q=ana&ordenar=si'); record('search_sort', [r['nombre'] for r in contexts[-1]['context']['contactos']])
    client.get('/principal?ordenar=si'); record('sort_all', [r['nombre'] for r in contexts[-1]['context']['contactos']])
    client.get('/estadisticas'); ctx=contexts[-1]['context']; record('stats',{'total':ctx['total'],'favorites':ctx['favorites'],'recent_names':[r['nombre'] for r in ctx['contactos_recientes']]})
    reset([]); client.get('/estadisticas'); ctx=contexts[-1]['context']; record('empty_stats',{'total':ctx['total'],'favorites':ctx['favorites']})
    reset([a]); response=client.post('/guardar_contacto',data={**valid,'nombre_completo':'<script>alert(1)</script>'})
    html=client.get('/principal').get_data(as_text=True)
    record('html_escape',{'escaped':'&lt;script&gt;alert(1)&lt;/script&gt;' in html,'raw':'<script>alert(1)</script>' in html})
    reset([a]); Path(module.EXCEL_FILE).write_text('not-an-xlsx')
    record('corrupt_workbook',client.get('/principal').status_code)
    Path(module.EXCEL_FILE).unlink()
    record('missing_workbook_detail',client.get('/mas_informacion/2').status_code)
    record('missing_workbook_list_recreates',client.get('/principal').status_code)
    record('missing_workbook_rows_after_list',len(rows())-1)
    # Controlled overlap: both requests load the same revision; B saves after A.
    reset([a]); original_loader=module.load_workbook
    barrier=threading.Barrier(2); saved_a=threading.Event()
    def overlapping_load(path):
        wb=original_loader(path); actual_save=wb.save
        barrier.wait(timeout=10)
        def save(path):
            if threading.current_thread().name=='request-B': saved_a.wait(timeout=10)
            actual_save(path)
            if threading.current_thread().name=='request-A': saved_a.set()
        wb.save=save
        return wb
    module.load_workbook=overlapping_load
    responses=[]
    def create(name):
        with app.test_client() as concurrent_client:
            res=concurrent_client.post('/guardar_contacto',data={**valid,'nombre_completo':name})
            responses.append({'name':name,'status':res.status_code})
    threads=[threading.Thread(target=create,args=('Primero Temporal',),name='request-A'),threading.Thread(target=create,args=('Segundo Temporal',),name='request-B')]
    for thread in threads: thread.start()
    for thread in threads: thread.join(timeout=15)
    module.load_workbook=original_loader
    record('concurrent_create_lost_update',{'responses':responses,'persisted_names':[r[0] for r in rows()[1:]]})
    record('runtime_dependencies',{n:version(n) for n in ['Flask','Werkzeug','openpyxl']})
    record('route_methods',{r.rule:sorted(r.methods) for r in app.url_map.iter_rules()})
    template_rendered.disconnect(capture,app)
    os.chdir(ROOT)

(ROOT/'integrity-test-results.json').write_text(json.dumps(results,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(results,ensure_ascii=True,indent=2))
